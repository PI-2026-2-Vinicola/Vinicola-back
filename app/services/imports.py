"""
Importação de dados reais (CSV, Excel .xlsx e JSON).

Fluxo: arquivo → leitura da tabela → normalização dos cabeçalhos → validação linha a linha
→ detecção de duplicados (no banco e no próprio arquivo) → pré-visualização ou gravação
em uma única transação → registro em import_jobs e na auditoria.
"""

from __future__ import annotations

import csv
import io
import json
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable

from sqlalchemy.orm import Session

from ..catalog import CLASSIFICATION_BY_QUALITY, MATURATION_LABEL, MATURATION_UNKNOWN, MATURATIONS, QUALITY_LABEL
from ..config import get_settings
from ..models import EnvironmentReading, Reading, Sensor, Variety
from ..schemas import SENSOR_ID_RE
from .readings import next_code
from .timeutil import as_utc, from_db, utcnow

KINDS = ("readings", "sensors", "environment")

TEMPLATES: dict[str, list[str]] = {
    "readings": ["sensor_id", "captured_at", "variety_id", "quality", "confidence", "maturation", "classification", "visual_condition", "observations", "clusters_detected", "model_version"],
    "sensors": ["id", "name", "block", "location", "latitude", "longitude", "variety_id", "device", "firmware", "capture_interval_min", "installed_at"],
    "environment": ["sensor_id", "measured_at", "temperature_c", "humidity_pct", "luminosity_lux", "soil_moisture_pct"],
}

REQUIRED: dict[str, list[str]] = {
    "readings": ["sensor_id", "captured_at", "quality", "confidence"],
    "sensors": ["id", "name", "block", "location", "variety_id"],
    "environment": ["sensor_id", "measured_at"],
}

ALIASES = {
    "sensor": "sensor_id", "id_sensor": "sensor_id", "sensor_codigo": "sensor_id",
    "data": "captured_at", "data_hora": "captured_at", "datahora": "captured_at", "capturado_em": "captured_at", "timestamp": "captured_at", "captured": "captured_at", "capturedat": "captured_at",
    "variedade": "variety_id", "variety": "variety_id", "uva": "variety_id", "varietyid": "variety_id",
    "qualidade": "quality", "confianca": "confidence", "confidencia": "confidence",
    "maturacao": "maturation", "estagio": "maturation", "estagio_maturacao": "maturation",
    "classificacao": "classification", "condicao_visual": "visual_condition", "condicao": "visual_condition", "visualcondition": "visual_condition",
    "observacoes": "observations", "observacao": "observations", "obs": "observations",
    "cachos": "clusters_detected", "cachos_detectados": "clusters_detected", "modelo": "model_version",
    "codigo": "id", "nome": "name", "bloco": "block", "talhao": "block", "localizacao": "location", "local": "location",
    "lat": "latitude", "lng": "longitude", "lon": "longitude", "long": "longitude",
    "dispositivo": "device", "intervalo": "capture_interval_min", "intervalo_captura": "capture_interval_min", "instalado_em": "installed_at", "instalacao": "installed_at",
    "medido_em": "measured_at", "data_medicao": "measured_at", "measuredat": "measured_at",
    "temperatura": "temperature_c", "temperatura_c": "temperature_c", "temp": "temperature_c",
    "umidade": "humidity_pct", "umidade_ar": "humidity_pct", "umidade_pct": "humidity_pct",
    "luminosidade": "luminosity_lux", "lux": "luminosity_lux",
    "umidade_solo": "soil_moisture_pct", "umidade_do_solo": "soil_moisture_pct",
}


class ImportError_(ValueError):
    """Erro que invalida o arquivo inteiro (formato, tamanho, cabeçalho)."""


def _norm(text: str) -> str:
    text = unicodedata.normalize("NFD", str(text)).encode("ascii", "ignore").decode()
    return text.strip().lower().replace(" ", "_").replace("-", "_").replace("(", "").replace(")", "").replace("%", "pct").replace("°", "")


def _header(name: str) -> str:
    key = _norm(name)
    return ALIASES.get(key, key)


def read_table(filename: str, data: bytes) -> list[dict[str, Any]]:
    settings = get_settings()
    if len(data) > settings.max_import_mb * 1_048_576:
        raise ImportError_(f"Arquivo maior que {settings.max_import_mb} MB.")
    if not data.strip():
        raise ImportError_("Arquivo vazio.")
    ext = Path(filename).suffix.lower()
    if ext == ".csv" or ext == ".txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        sample = text[:4096]
        delimiter = max([";", ",", "\t"], key=sample.count)
        rows = list(csv.DictReader(io.StringIO(text), delimiter=delimiter))
    elif ext == ".xlsx":
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ImportError_("Não foi possível ler a planilha. Envie um arquivo .xlsx válido.") from exc
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            raise ImportError_("A primeira linha da planilha deve conter os nomes das colunas.")
        names = [str(h) if h is not None else "" for h in header]
        rows = [dict(zip(names, r)) for r in it if any(v not in (None, "") for v in r)]
    elif ext == ".json":
        try:
            parsed = json.loads(data.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ImportError_(f"JSON inválido: {exc}") from exc
        if isinstance(parsed, dict):
            parsed = parsed.get("items") or parsed.get("data") or parsed.get("rows")
        if not isinstance(parsed, list) or not all(isinstance(r, dict) for r in parsed):
            raise ImportError_("O JSON deve ser uma lista de objetos (ou um objeto com a chave \"items\").")
        rows = parsed
    else:
        raise ImportError_("Formato não suportado. Use .csv, .xlsx ou .json.")
    if len(rows) > settings.max_import_rows:
        raise ImportError_(f"O arquivo tem {len(rows)} linhas; o limite por importação é {settings.max_import_rows}.")
    return [{_header(k): v for k, v in r.items() if k is not None and str(k).strip()} for r in rows]


# ----------------------------------------------------------------------------- conversões
class FieldError(ValueError):
    def __init__(self, field: str, message: str):
        super().__init__(message)
        self.field = field


def _empty(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _text(row: dict, key: str, max_len: int, required: bool = False) -> str | None:
    v = row.get(key)
    if _empty(v):
        if required:
            raise FieldError(key, "campo obrigatório")
        return None
    s = str(v).strip()
    if len(s) > max_len:
        raise FieldError(key, f"máximo de {max_len} caracteres")
    return s


def _number(row: dict, key: str, lo: float | None = None, hi: float | None = None) -> float | None:
    v = row.get(key)
    if _empty(v):
        return None
    try:
        n = float(str(v).strip().replace(",", ".")) if not isinstance(v, (int, float)) else float(v)
    except ValueError:
        raise FieldError(key, f"“{v}” não é um número")
    if (lo is not None and n < lo) or (hi is not None and n > hi):
        raise FieldError(key, f"valor {n:g} fora do intervalo permitido ({lo:g} a {hi:g})")
    return n


DATETIME_FORMATS = ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d/%m/%Y", "%Y-%m-%d")


def _datetime(row: dict, key: str) -> datetime:
    v = row.get(key)
    if _empty(v):
        raise FieldError(key, "campo obrigatório")
    if isinstance(v, datetime):
        dt = v
    elif isinstance(v, date):
        dt = datetime.combine(v, datetime.min.time())
    else:
        s = str(v).strip()
        dt = None
        try:
            dt = datetime.fromisoformat(s.replace("Z", "+00:00"))
        except ValueError:
            for fmt in DATETIME_FORMATS:
                try:
                    dt = datetime.strptime(s, fmt)
                    break
                except ValueError:
                    continue
        if dt is None:
            raise FieldError(key, f"data/hora “{s}” inválida (use AAAA-MM-DD HH:MM ou DD/MM/AAAA HH:MM)")
    dt = as_utc(dt)
    if dt > utcnow() + timedelta(minutes=5):
        raise FieldError(key, "data no futuro")
    if dt.year < 2000:
        raise FieldError(key, "data anterior a 2000")
    return dt


def _date(row: dict, key: str) -> date | None:
    v = row.get(key)
    if _empty(v):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise FieldError(key, f"data “{s}” inválida")


QUALITY_ALIASES = {_norm(k): k for k in QUALITY_LABEL} | {_norm(v): k for k, v in QUALITY_LABEL.items()} | {"critico": "critica", "ruim": "critica", "atencao": "atencao"}
CLASS_ALIASES = {_norm(c): c for c in CLASSIFICATION_BY_QUALITY.values()}
MATURATION_ALIASES = {_norm(k): k for k in MATURATION_LABEL} | {_norm(v): k for k, v in MATURATION_LABEL.items()} | {"pintor_veraison": "pintor", "veraison": "pintor"}


@dataclass
class Context:
    db: Session
    sensors: dict[str, Sensor]
    varieties: dict[str, str]  # chave normalizada (id ou nome) → id
    seen: set[tuple] = field(default_factory=set)


def _variety(ctx: Context, row: dict, key: str = "variety_id", required: bool = False) -> str | None:
    v = row.get(key)
    if _empty(v):
        if required:
            raise FieldError(key, "campo obrigatório")
        return None
    vid = ctx.varieties.get(_norm(v).replace("_", "-")) or ctx.varieties.get(_norm(v))
    if not vid:
        raise FieldError(key, f"variedade “{v}” não cadastrada")
    return vid


def _sensor(ctx: Context, row: dict) -> Sensor:
    sid = _text(row, "sensor_id", 20, required=True).upper()  # type: ignore[union-attr]
    sensor = ctx.sensors.get(sid)
    if not sensor:
        raise FieldError("sensor_id", f"sensor “{sid}” não cadastrado")
    return sensor


def validate_reading(ctx: Context, row: dict) -> tuple[tuple, dict]:
    sensor = _sensor(ctx, row)
    captured = _datetime(row, "captured_at")
    variety = _variety(ctx, row) or sensor.variety_id
    q_raw = row.get("quality")
    if _empty(q_raw):
        raise FieldError("quality", "campo obrigatório")
    quality = QUALITY_ALIASES.get(_norm(q_raw))
    if not quality:
        raise FieldError("quality", f"qualidade “{q_raw}” inválida (use Boa, Atenção ou Necessita atenção)")
    confidence = _number(row, "confidence", 0, 100)
    if confidence is None:
        raise FieldError("confidence", "campo obrigatório")
    if confidence > 1:
        confidence = confidence / 100
    classification = CLASSIFICATION_BY_QUALITY[quality]
    c_raw = row.get("classification")
    if not _empty(c_raw):
        given = CLASS_ALIASES.get(_norm(c_raw))
        if not given:
            raise FieldError("classification", f"classificação “{c_raw}” inválida")
        if given != classification:
            raise FieldError("classification", f"“{c_raw}” não corresponde à qualidade “{QUALITY_LABEL[quality]}”")
    m_raw = row.get("maturation")
    maturation = MATURATION_UNKNOWN
    if not _empty(m_raw):
        maturation = MATURATION_ALIASES.get(_norm(m_raw)) or ""
        if maturation not in MATURATIONS:
            raise FieldError("maturation", f"estágio “{m_raw}” inválido")
    clusters = _number(row, "clusters_detected", 0, 50)
    record = {
        "sensor_id": sensor.id,
        "captured_at": captured,
        "variety_id": variety,
        "quality": quality,
        "confidence": round(confidence, 4),
        "maturation": maturation,
        "classification": classification,
        "visual_condition": _text(row, "visual_condition", 60) or QUALITY_LABEL[quality],
        "observations": _text(row, "observations", 2000) or "",
        "clusters_detected": int(clusters) if clusters is not None else 1,
        "model_version": _text(row, "model_version", 80) or "Importado",
    }
    return (sensor.id, captured.replace(microsecond=0)), record


def validate_sensor(ctx: Context, row: dict) -> tuple[tuple, dict]:
    sid = _text(row, "id", 20, required=True).upper()  # type: ignore[union-attr]
    if not SENSOR_ID_RE.match(sid):
        raise FieldError("id", "use 2 a 20 caracteres: letras, números, '-' ou '_'")
    lat = _number(row, "latitude", -90, 90)
    lng = _number(row, "longitude", -180, 180)
    if (lat is None) != (lng is None):
        raise FieldError("latitude", "informe latitude e longitude juntas")
    if lat == 0 and lng == 0:
        raise FieldError("latitude", "coordenadas 0,0 não são válidas")
    interval = _number(row, "capture_interval_min", 5, 1440)
    record = {
        "id": sid,
        "name": _text(row, "name", 80, required=True),
        "block": _text(row, "block", 40, required=True),
        "location": _text(row, "location", 120, required=True),
        "latitude": lat,
        "longitude": lng,
        "variety_id": _variety(ctx, row, required=True),
        "device": _text(row, "device", 60),
        "firmware": _text(row, "firmware", 20),
        "capture_interval_min": int(interval) if interval else 90,
        "installed_at": _date(row, "installed_at"),
    }
    return (sid,), record


def validate_environment(ctx: Context, row: dict) -> tuple[tuple, dict]:
    sensor = _sensor(ctx, row)
    measured = _datetime(row, "measured_at")
    record = {
        "sensor_id": sensor.id,
        "measured_at": measured,
        "temperature_c": _number(row, "temperature_c", -20, 60),
        "humidity_pct": _number(row, "humidity_pct", 0, 100),
        "luminosity_lux": _number(row, "luminosity_lux", 0, 200000),
        "soil_moisture_pct": _number(row, "soil_moisture_pct", 0, 100),
    }
    if all(record[k] is None for k in ("temperature_c", "humidity_pct", "luminosity_lux", "soil_moisture_pct")):
        raise FieldError("temperature_c", "informe ao menos uma medição (temperatura, umidade, luminosidade ou umidade do solo)")
    return (sensor.id, measured.replace(microsecond=0)), record


VALIDATORS: dict[str, Callable[[Context, dict], tuple[tuple, dict]]] = {
    "readings": validate_reading,
    "sensors": validate_sensor,
    "environment": validate_environment,
}


def _context(db: Session) -> Context:
    sensors = {s.id: s for s in db.query(Sensor).all()}
    varieties: dict[str, str] = {}
    for v in db.query(Variety).all():
        varieties[v.id] = v.id
        varieties[_norm(v.name).replace("_", "-")] = v.id
        varieties[_norm(v.name)] = v.id
    return Context(db, sensors, varieties)


def _existing_keys(db: Session, kind: str, keys: list[tuple]) -> set[tuple]:
    """Busca no banco apenas os registros no intervalo de datas do arquivo (evita varrer a tabela toda)."""
    if not keys:
        return set()
    if kind == "sensors":
        ids = [k[0] for k in keys]
        return {(sid,) for (sid,) in db.query(Sensor.id).filter(Sensor.id.in_(ids))}
    model, col = (Reading, Reading.captured_at) if kind == "readings" else (EnvironmentReading, EnvironmentReading.measured_at)
    sensors = sorted({k[0] for k in keys})
    lo = min(k[1] for k in keys)
    hi = max(k[1] for k in keys) + timedelta(seconds=1)
    rows = db.query(model.sensor_id, col).filter(model.sensor_id.in_(sensors), col >= lo, col < hi)
    return {(sid, from_db(t).replace(microsecond=0)) for sid, t in rows}


@dataclass
class ImportResult:
    kind: str
    filename: str
    total: int = 0
    valid: list[tuple[tuple, dict]] = field(default_factory=list)
    duplicates: list[tuple[tuple, dict]] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    invalid_rows: int = 0
    columns: list[str] = field(default_factory=list)


def analyze_file(db: Session, kind: str, filename: str, data: bytes) -> ImportResult:
    if kind not in KINDS:
        raise ImportError_("Tipo de importação inválido.")
    rows = read_table(filename, data)
    result = ImportResult(kind, filename, total=len(rows))
    if not rows:
        raise ImportError_("Nenhuma linha de dados encontrada.")
    columns = sorted({k for r in rows for k in r})
    result.columns = columns
    missing = [c for c in REQUIRED[kind] if c not in columns]
    if missing:
        raise ImportError_(f"Colunas obrigatórias ausentes: {', '.join(missing)}. Baixe o modelo para ver o formato esperado.")
    ctx = _context(db)
    validator = VALIDATORS[kind]
    accepted: list[tuple[tuple, dict]] = []
    for i, row in enumerate(rows, start=2):  # linha 1 = cabeçalho
        try:
            key, record = validator(ctx, row)
        except FieldError as exc:
            result.invalid_rows += 1
            if len(result.errors) < 500:
                result.errors.append({"row": i, "field": exc.field, "message": str(exc)})
            continue
        if key in ctx.seen:
            result.invalid_rows += 1
            if len(result.errors) < 500:
                result.errors.append({"row": i, "field": None, "message": "linha repetida dentro do próprio arquivo"})
            continue
        ctx.seen.add(key)
        accepted.append((key, record))
    existing = _existing_keys(db, kind, [k for k, _ in accepted])
    for key, record in accepted:
        (result.duplicates if key in existing else result.valid).append((key, record))
    return result


def sample(records: list[tuple[tuple, dict]], n: int = 20) -> list[dict]:
    out = []
    for _, r in records[:n]:
        out.append({k: (v.isoformat() if isinstance(v, (datetime, date)) else v) for k, v in r.items()})
    return out


def persist(db: Session, result: ImportResult, on_duplicate: str, user_id: int | None) -> tuple[int, int]:
    """Grava as linhas válidas (e atualiza duplicados, se pedido). Retorna (inseridas, atualizadas)."""
    inserted = updated = 0
    if result.kind == "readings":
        new = []
        for _, r in result.valid:
            reading = Reading(**r, source="importacao", processing_ms=0, stage="concluida", created_by=user_id)
            db.add(reading)
            new.append(reading)
        db.flush()
        for reading in new:
            reading.code = next_code(reading)
        inserted = len(new)
        if on_duplicate == "update":
            for (sid, t), r in result.duplicates:
                existing = db.query(Reading).filter(Reading.sensor_id == sid, Reading.captured_at == t).first() or db.query(Reading).filter(
                    Reading.sensor_id == sid, Reading.captured_at >= t, Reading.captured_at < t + timedelta(seconds=1)
                ).first()
                if existing:
                    for k, v in r.items():
                        if k not in ("sensor_id", "captured_at"):
                            setattr(existing, k, v)
                    updated += 1
    elif result.kind == "sensors":
        for _, r in result.valid:
            db.add(Sensor(**r, active=True))
            inserted += 1
        if on_duplicate == "update":
            for (sid,), r in result.duplicates:
                s = db.get(Sensor, sid)
                if s:
                    for k, v in r.items():
                        if k != "id" and v is not None:
                            setattr(s, k, v)
                    updated += 1
    else:
        for _, r in result.valid:
            db.add(EnvironmentReading(**r, source="importacao"))
            inserted += 1
        if on_duplicate == "update":
            for (sid, t), r in result.duplicates:
                e = db.query(EnvironmentReading).filter(
                    EnvironmentReading.sensor_id == sid, EnvironmentReading.measured_at >= t, EnvironmentReading.measured_at < t + timedelta(seconds=1)
                ).first()
                if e:
                    for k in ("temperature_c", "humidity_pct", "luminosity_lux", "soil_moisture_pct"):
                        if r[k] is not None:
                            setattr(e, k, r[k])
                    updated += 1
    return inserted, updated


def template_csv(kind: str) -> str:
    return ";".join(TEMPLATES[kind]) + "\n"
